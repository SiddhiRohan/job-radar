"""Parse a base resume .docx into blocks, get a tailoring plan from Claude, apply it without touching formatting."""
import copy
import re

from docx import Document

import llm

HEADINGS = ("PROFESSIONAL SUMMARY", "WORK EXPERIENCE", "RESEARCH PUBLICATIONS", "SKILLS", "EDUCATION", "PROJECTS")
DATE = re.compile(r"(19|20)\d\d\s*$|Present\s*$")


def marked_text(p):
    """Paragraph text with **...** around bold runs, so inline bold survives a rewrite."""
    out, bold = "", False
    for r in p.runs:
        if bool(r.bold) != bold:
            out, bold = out + "**", not bold
        out += r.text
    return (out + ("**" if bold else "")).replace("****", "")


def parse(doc):
    """Locate summary, job blocks (header + bullets), skills line, coursework line by paragraph index."""
    ps, section, info = doc.paragraphs, None, {"summary": None, "jobs": [], "skills": None, "coursework": None}
    for i, p in enumerate(ps):
        t = p.text.strip()
        if not t:
            continue
        if t.rstrip("\t").upper() in HEADINGS:
            section = t.rstrip("\t").upper()
            continue
        if section == "PROFESSIONAL SUMMARY" and info["summary"] is None:
            info["summary"] = i
        elif section == "WORK EXPERIENCE":
            if "|" in t and DATE.search(t) and any(r.bold for r in p.runs):
                info["jobs"].append({"header": i, "bullets": []})
            elif info["jobs"]:
                info["jobs"][-1]["bullets"].append(i)
        elif section == "SKILLS" and info["skills"] is None:
            info["skills"] = i
        elif section == "EDUCATION" and t.lower().startswith("coursework"):
            info["coursework"] = i
    return info


def set_text(p, marked):
    """Replace a paragraph's text, keeping its style and the first run's font; **segments** become bold."""
    template = copy.deepcopy(p.runs[0]._r) if p.runs else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for k, seg in enumerate(re.split(r"\*\*", marked)):
        if not seg:
            continue
        run = p.add_run(seg)
        if template is not None:
            rpr = template.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
        run.bold = bool(k % 2)


def reorder(doc, job, order):
    """Move bullet paragraphs so they follow the header in `order` (indices into job['bullets'])."""
    ps = doc.paragraphs
    anchor = ps[job["header"]]._p
    for idx in order:
        el = ps[job["bullets"][idx]]._p
        anchor.addnext(el)
        anchor = el


PLAN_SCHEMA = {"type": "object", "additionalProperties": False, "properties": {
    "headline": {"type": "string", "enum": ["Data Engineer", "Data Scientist", "ML Engineer", "AI Engineer"]},
    "summary": {"type": "string"},
    "jobs": {"type": "array", "items": {"type": "object", "additionalProperties": False, "properties": {
        "job": {"type": "integer"}, "order": {"type": "array", "items": {"type": "integer"}},
        "rewrites": {"type": "array", "items": {"type": "object", "additionalProperties": False, "properties": {
            "bullet": {"type": "integer"}, "text": {"type": "string"}, "reason": {"type": "string"}},
            "required": ["bullet", "text", "reason"]}}}, "required": ["job", "order", "rewrites"]}},
    "skills": {"type": "string"}, "coursework": {"type": "string"},
    "hard_to_defend": {"type": "array", "items": {"type": "string"}},
    "questions_for_rohan": {"type": "array", "items": {"type": "string"}},
    "fit_notes": {"type": "string"}},
    "required": ["headline", "summary", "jobs", "skills", "coursework", "hard_to_defend", "questions_for_rohan", "fit_notes"]}

PLAN_RULES = """You tailor a resume to one job posting. LOCKED: header, employer names, titles, dates, GPA, every number.
Allowed: rewrite the summary to mirror the posting's phrasing; reorder bullets within a job by relevance;
lightly reword bullets to mirror the posting's terms WITHOUT adding tools, scope, or facts not already in the base;
adjust the Skills line (only tools already on the base resume, reordered/pruned); pick coursework already listed.
Keep **bold** markers on methods, domain terms, and headline metrics; use ** in your text the same way.
Never use em dashes. If a bullet would need something new to match the posting, put a question in questions_for_rohan
instead of writing it. Flag bullets that would be hard to defend in an interview. Return only text for bullets you
actually change; keep rewrites minimal. Headline rotates by role: Data Engineer, Data Scientist, ML Engineer, AI Engineer."""


def make_plan(doc, info, jd_text, job_meta, profile):
    ps = doc.paragraphs
    blocks = [f"SUMMARY: {marked_text(ps[info['summary']])}"]
    for k, j in enumerate(info["jobs"]):
        blocks.append(f"JOB {k}: {ps[j['header']].text.strip()}")
        blocks += [f"  bullet {b}: {marked_text(ps[i])}" for b, i in enumerate(j["bullets"])]
    blocks.append(f"SKILLS: {ps[info['skills']].text.strip()}")
    if info["coursework"] is not None:
        blocks.append(f"COURSEWORK: {ps[info['coursework']].text.strip()}")
    user = (f"POSTING: {job_meta['title']} at {job_meta['company']}\n{jd_text[:12000]}\n\nBASE RESUME:\n"
            + "\n".join(blocks) + "\n\nReturn the tailoring plan.")
    return llm.complete(PLAN_RULES + "\n\nCANDIDATE PROFILE:\n" + profile, user, PLAN_SCHEMA, max_tokens=6000)[0]


def apply_plan(doc, info, plan):
    """Apply the plan; return [(where, before, after)] for the terminal diff and notes.md."""
    ps, changes = doc.paragraphs, []
    def change(where, idx, new):
        old = marked_text(ps[idx])
        if new and new.strip() != old.strip():
            set_text(ps[idx], new)
            changes.append((where, old, new))
    change("summary", info["summary"], plan["summary"])
    for pj in plan["jobs"]:
        job = info["jobs"][pj["job"]]
        for rw in pj["rewrites"]:
            if 0 <= rw["bullet"] < len(job["bullets"]):
                change(f"job {pj['job']} bullet {rw['bullet']}", job["bullets"][rw["bullet"]], rw["text"])
        order = [i for i in pj["order"] if 0 <= i < len(job["bullets"])]
        if order and order != sorted(order) and len(set(order)) == len(job["bullets"]):
            changes.append((f"job {pj['job']} order", " ".join(map(str, range(len(order)))), " ".join(map(str, order))))
            reorder(doc, job, order)
    change("skills", info["skills"], plan["skills"])
    if info["coursework"] is not None and plan["coursework"]:
        change("coursework", info["coursework"], plan["coursework"])
    return changes
